"""Tambahkan master toko dengan kode zona deterministik berdasarkan pola UJP.

Lokal:
    python scripts/penambahan_master_toko.py data_master.xlsx -o hasil_master_toko.xlsx

Google Colab:
    %run scripts/penambahan_master_toko.py

Workbook sumber tidak ditimpa. Format sheet ``DATA MASTER`` dipertahankan,
kemudian kolom ``KODE ZONA`` dan ``NAMA ZONA`` diisi. Sheet mapping dan audit
ditambahkan untuk memudahkan pemeriksaan hasil.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import unicodedata
from copy import copy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from numbers import Real
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from google.colab import files
except ImportError:  # Memungkinkan fungsi diuji di luar Google Colab.
    files = None


# =========================
# KONFIGURASI
# =========================
SOURCE_SHEET = "DATA MASTER"
OP_COLUMN = "OP"
KECAMATAN_COLUMN = "KECAMATAN"
ZONE_CODE_COLUMN = "KODE ZONA"
ZONE_NAME_COLUMN = "NAMA ZONA"
MAPPING_SHEET = "MAPPING_ZONA"
AUDIT_SHEET = "AUDIT_ZONA"
MAX_ZONE_CODE_LENGTH = 15
UJP_DECIMAL_PLACES = 2
EMPTY_UJP_AS_ZERO = True
AUTO_DETECT_EXTRA_UJP_COLUMNS = True
UJP_COLUMN_PREFIX = "UJP "

# Daftar standar dan urutan prioritas kolom UJP.
UJP_COLUMNS = [
    "UJP CDD",
    "UJP CDD-L",
    "UJP CDD-C",
    "UJP CDE",
    "UJP CDE-L",
    "UJP CDE-C",
    "UJP L300",
    "UJP FUSO",
    "UJP CDD-LC",
    "UJP CDE-LC",
    "UJP Tronton",
    "UJP Wingbox",
    "UJP Traga",
]

DIRECTION_ABBREVIATIONS = {
    "BARAT": "BAR",
    "TIMUR": "TIM",
    "UTARA": "UTA",
    "SELATAN": "SEL",
    "TENGAH": "TEN",
}
ADMINISTRATIVE_WORDS = {"KECAMATAN", "KEC"}


def normalize_header(value) -> str:
    """Menormalkan nama header tanpa mengubah header asli di workbook."""
    return re.sub(r"\s+", " ", str(value).strip()).upper()


def is_blank(value) -> bool:
    if value is None:
        return True
    try:
        if bool(pd.isna(value)):
            return True
    except (TypeError, ValueError):
        pass
    return str(value).strip() == ""


def normalized_words(value) -> list[str]:
    """Menghasilkan token A-Z/0-9 yang konsisten dari teks OP/kecamatan."""
    if is_blank(value):
        return []
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    tokens = re.findall(r"[A-Z0-9]+", text.upper())
    tokens = [token for token in tokens if token not in ADMINISTRATIVE_WORDS]
    return tokens


def normalize_op(value) -> str:
    return "".join(normalized_words(value))


def normalize_kecamatan(value) -> str:
    return " ".join(normalized_words(value))


def zone_name_stem(value) -> str:
    return "".join(normalized_words(value))


def parse_numeric_text(value: str) -> Decimal | None:
    """Mencoba membaca teks angka format umum Indonesia maupun internasional."""
    text = value.strip().upper().replace("RP", "").replace(" ", "")
    if not text:
        return None
    if not re.fullmatch(r"[-+]?\d[\d.,]*", text):
        return None
    if "." in text and "," in text:
        # Pemisah paling kanan dianggap sebagai pemisah desimal.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        right = text.rsplit(",", 1)[1]
        text = text.replace(",", ".") if len(right) <= 2 else text.replace(",", "")
    elif "." in text:
        right = text.rsplit(".", 1)[1]
        # 1.000 lazim berarti seribu; 1000.50 berarti angka desimal.
        if len(right) == 3:
            text = text.replace(".", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def canonical_ujp(value) -> str:
    """Menyamakan 1000, 1000.0, dan '1.000' agar tidak dianggap pola berbeda."""
    if is_blank(value):
        return "N:0" if EMPTY_UJP_AS_ZERO else "EMPTY"
    decimal_value = None
    if isinstance(value, (Real, Decimal)) and not isinstance(value, bool):
        if not math.isfinite(float(value)):
            return "N:0" if EMPTY_UJP_AS_ZERO else "EMPTY"
        decimal_value = Decimal(str(value))
    else:
        decimal_value = parse_numeric_text(str(value))
    if decimal_value is not None:
        quantizer = Decimal("1").scaleb(-UJP_DECIMAL_PLACES)
        decimal_value = decimal_value.quantize(quantizer, rounding=ROUND_HALF_UP)
        if decimal_value == 0:
            decimal_value = Decimal(0)
        normalized = format(decimal_value.normalize(), "f")
        return f"N:{normalized}"
    normalized_text = re.sub(r"\s+", " ", str(value).strip()).upper()
    return f"T:{normalized_text}"


def detect_ujp_columns(dataframe_columns: list[str]) -> tuple[list[str], list[str], list[str]]:
    """
    Menentukan kolom UJP yang ikut membentuk pola tarif.
    Return:
    - available: kolom standar yang tersedia + kolom UJP tambahan terdeteksi.
    - missing_standard: kolom standar yang belum ada di file (bukan error).
    - auto_detected: kolom UJP baru di luar daftar standar.
    """
    existing_columns = list(dataframe_columns)
    standard_columns = [normalize_header(column) for column in UJP_COLUMNS]
    available_standard = [
        column for column in standard_columns if column in existing_columns
    ]
    missing_standard = [
        column for column in standard_columns if column not in existing_columns
    ]
    auto_detected = []
    if AUTO_DETECT_EXTRA_UJP_COLUMNS:
        prefix = normalize_header(UJP_COLUMN_PREFIX)
        auto_detected = sorted(
            column
            for column in existing_columns
            if column.startswith(prefix)
            and column not in standard_columns
            and column not in available_standard
        )
    available = available_standard + auto_detected
    return available, missing_standard, auto_detected


def semantic_kecamatan_code(kecamatan_norm: str, max_length: int) -> str:
    """Membuat singkatan yang mudah dibaca, tidak terpaku pada 4 awal + 3 akhir."""
    if max_length < 1:
        raise ValueError("Ruang kode kecamatan kurang dari 1 karakter.")
    words = kecamatan_norm.split()
    joined = "".join(words)
    if len(joined) <= max_length:
        return joined
    if len(words) == 1:
        return joined[:max_length]
    # Contoh: BLORA KOTA TIMUR -> BLO + TIM = BLOTIM.
    if words[-1] in DIRECTION_ABBREVIATIONS and max_length >= 4:
        tail = DIRECTION_ABBREVIATIONS[words[-1]][: min(3, max_length - 1)]
        head_length = max_length - len(tail)
        return words[0][:head_length] + tail
    # Contoh: BLORA KOTA -> BLO + KOT = BLOKOT.
    left_length = (max_length + 1) // 2
    right_length = max_length - left_length
    return words[0][:left_length] + words[-1][:right_length]


def collision_kecamatan_code(kecamatan_norm: str, max_length: int) -> str:
    """Alternatif yang lebih informatif saat kandidat utama saling bertabrakan."""
    words = kecamatan_norm.split()
    joined = "".join(words)
    if len(joined) <= max_length:
        return joined
    if len(words) == 1:
        # KEDUNGJATI -> KEDATI; KEDUNGPRING -> KEDING.
        left_length = (max_length + 1) // 2
        right_length = max_length - left_length
        return joined[:left_length] + joined[-right_length:]
    if words[-1] in DIRECTION_ABBREVIATIONS and max_length >= 5:
        tail = DIRECTION_ABBREVIATIONS[words[-1]][:3]
        head_space = max_length - len(tail)
        if len(words) >= 3 and head_space >= 2:
            # Jika BLO+TIM bentrok, sertakan inisial kata tengah: BLK+TIM / BLD+TIM.
            head = words[0][: head_space - 1] + words[-2][:1]
        else:
            head = words[0][:head_space]
        return (head + tail)[:max_length]
    left_length = (max_length + 1) // 2
    right_length = max_length - left_length
    return joined[:left_length] + joined[-right_length:]


def alpha_hash(text: str, length: int, salt: int = 0) -> str:
    """Hash hanya huruf A-Z agar batas antara kode kecamatan dan suffix tetap jelas."""
    digest = hashlib.sha256(f"{text}|{salt}".encode("utf-8")).digest()
    number = int.from_bytes(digest, "big")
    chars = []
    for _ in range(length):
        number, remainder = divmod(number, 26)
        chars.append(chr(ord("A") + remainder))
    return "".join(chars)


def unique_kecamatan_codes(items: list[dict]) -> dict[tuple[str, str], str]:
    """Menjamin singkatan kecamatan unik pada setiap OP."""
    result = {}
    by_op = {}
    for item in items:
        by_op.setdefault(item["op"], []).append(item)
    for op, op_items in sorted(by_op.items()):
        ordered = sorted(op_items, key=lambda item: item["kecamatan"])
        initial = {
            item["kecamatan"]: semantic_kecamatan_code(
                item["kecamatan"], item["max_kec_length"]
            )
            for item in ordered
        }
        grouped = {}
        for kecamatan, code in initial.items():
            grouped.setdefault(code, []).append(kecamatan)
        used_codes = set()
        # Kode yang sejak awal tidak bertabrakan dipertahankan.
        for code, kecamatan_list in sorted(grouped.items()):
            if len(kecamatan_list) == 1:
                kecamatan = kecamatan_list[0]
                result[(op, kecamatan)] = code
                used_codes.add(code)
        # Pada kelompok bentrok, coba kandidat alternatif yang tetap mudah dibaca.
        for code, kecamatan_list in sorted(grouped.items()):
            if len(kecamatan_list) <= 1:
                continue
            for kecamatan in sorted(kecamatan_list):
                item = next(x for x in ordered if x["kecamatan"] == kecamatan)
                max_length = item["max_kec_length"]
                readable_alternative = collision_kecamatan_code(kecamatan, max_length)
                candidate = (
                    readable_alternative
                    if readable_alternative not in used_codes
                    else None
                )
                if candidate is None:
                    for hash_length in range(1, max_length + 1):
                        prefix_length = max_length - hash_length
                        prefix = readable_alternative[:prefix_length]
                        for salt in range(1000):
                            proposed = prefix + alpha_hash(kecamatan, hash_length, salt)
                            if proposed not in used_codes:
                                candidate = proposed
                                break
                        if candidate is not None:
                            break
                if candidate is None:
                    raise ValueError(
                        f"Tidak dapat membuat kode kecamatan unik untuk OP {op}: {kecamatan}."
                    )
                result[(op, kecamatan)] = candidate
                used_codes.add(candidate)
    return result


def build_zone_mapping(
    df: pd.DataFrame, available_ujp_columns: list[str]
) -> tuple[pd.DataFrame, pd.DataFrame]:
    work = df.copy()
    work["__OP"] = work[OP_COLUMN].map(normalize_op)
    work["__KECAMATAN"] = work[KECAMATAN_COLUMN].map(normalize_kecamatan)
    work["__NAMA_ZONA"] = work[KECAMATAN_COLUMN].map(zone_name_stem)
    invalid_op = work.index[work["__OP"].eq("")].tolist()
    invalid_kec = work.index[work["__KECAMATAN"].eq("")].tolist()
    if invalid_op or invalid_kec:
        messages = []
        if invalid_op:
            messages.append(f"OP kosong/tidak valid pada baris Excel: {[i + 2 for i in invalid_op[:20]]}")
        if invalid_kec:
            messages.append(
                f"KECAMATAN kosong/tidak valid pada baris Excel: {[i + 2 for i in invalid_kec[:20]]}"
            )
        raise ValueError("; ".join(messages))
    for column in available_ujp_columns:
        work[f"__{column}"] = work[column].map(canonical_ujp)
    def pattern_from_row(row) -> tuple[str, ...]:
        return tuple(row[f"__{column}"] for column in available_ujp_columns)
    work["__PATTERN"] = work.apply(pattern_from_row, axis=1)
    unique_zone_keys = work[
        ["__OP", "__KECAMATAN", "__NAMA_ZONA", "__PATTERN"]
    ].drop_duplicates()
    pattern_counts = (
        unique_zone_keys.groupby(["__OP", "__KECAMATAN"])["__PATTERN"]
        .nunique()
        .to_dict()
    )
    district_items = []
    for op, kecamatan in sorted(pattern_counts):
        count = pattern_counts[(op, kecamatan)]
        suffix_digits = len(str(count))
        max_kec_length = MAX_ZONE_CODE_LENGTH - len(op) - suffix_digits
        if max_kec_length < 1:
            raise ValueError(
                f"OP '{op}' terlalu panjang untuk batas {MAX_ZONE_CODE_LENGTH} karakter "
                f"dan {count} pola UJP pada kecamatan '{kecamatan}'."
            )
        district_items.append(
            {
                "op": op,
                "kecamatan": kecamatan,
                "pattern_count": count,
                "max_kec_length": max_kec_length,
            }
        )
    kecamatan_code_map = unique_kecamatan_codes(district_items)
    suffix_map = {}
    for (op, kecamatan), group in unique_zone_keys.groupby(["__OP", "__KECAMATAN"]):
        patterns = sorted(group["__PATTERN"].tolist())
        for suffix, pattern in enumerate(patterns, start=1):
            suffix_map[(op, kecamatan, pattern)] = suffix
    records = []
    for _, row in unique_zone_keys.iterrows():
        op = row["__OP"]
        kecamatan = row["__KECAMATAN"]
        pattern = row["__PATTERN"]
        suffix = suffix_map[(op, kecamatan, pattern)]
        kec_code = kecamatan_code_map[(op, kecamatan)]
        zone_code = f"{op}{kec_code}{suffix}"
        zone_name = f"{row['__NAMA_ZONA']}{suffix}"
        pattern_payload = {
            column: value for column, value in zip(available_ujp_columns, pattern)
        }
        record = {
            "OP": op,
            "KECAMATAN": kecamatan,
            "KODE_KECAMATAN": kec_code,
            "SUFFIX_UJP": suffix,
            "KODE_ZONA": zone_code,
            "NAMA_ZONA": zone_name,
            "UJP_SIGNATURE": json.dumps(
                pattern_payload, ensure_ascii=False, separators=(",", ":")
            ),
            "__PATTERN": pattern,
        }
        record.update({f"POLA {column}": value for column, value in pattern_payload.items()})
        records.append(record)
    mapping = pd.DataFrame(records).sort_values(
        ["OP", "KECAMATAN", "SUFFIX_UJP"], kind="stable"
    ).reset_index(drop=True)
    key_to_code = {
        (row["OP"], row["KECAMATAN"], row["__PATTERN"]): (
            row["KODE_ZONA"],
            row["NAMA_ZONA"],
        )
        for _, row in mapping.iterrows()
    }
    output_values = [
        key_to_code[(row["__OP"], row["__KECAMATAN"], row["__PATTERN"])]
        for _, row in work.iterrows()
    ]
    work[ZONE_CODE_COLUMN] = [value[0] for value in output_values]
    work[ZONE_NAME_COLUMN] = [value[1] for value in output_values]
    return work, mapping


def validate_result(work: pd.DataFrame, mapping: pd.DataFrame) -> pd.DataFrame:
    validations = []
    max_length = int(work[ZONE_CODE_COLUMN].str.len().max()) if len(work) else 0
    validations.append(
        {
            "VALIDASI": "Panjang kode zona maksimal 15 karakter",
            "STATUS": "OK" if max_length <= MAX_ZONE_CODE_LENGTH else "GAGAL",
            "DETAIL": f"Panjang maksimum aktual: {max_length}",
        }
    )
    invalid_format = work[
        ~work[ZONE_CODE_COLUMN].str.fullmatch(r"[A-Z0-9]+", na=False)
    ]
    validations.append(
        {
            "VALIDASI": "Kode zona hanya berisi A-Z dan 0-9",
            "STATUS": "OK" if invalid_format.empty else "GAGAL",
            "DETAIL": f"Baris tidak valid: {len(invalid_format)}",
        }
    )
    duplicate_zone = mapping[mapping.duplicated("KODE_ZONA", keep=False)]
    validations.append(
        {
            "VALIDASI": "Satu kode zona tidak dipakai kombinasi berbeda",
            "STATUS": "OK" if duplicate_zone.empty else "GAGAL",
            "DETAIL": f"Mapping bertabrakan: {len(duplicate_zone)}",
        }
    )
    district_collision = (
        mapping[["OP", "KECAMATAN", "KODE_KECAMATAN"]]
        .drop_duplicates()
        .groupby(["OP", "KODE_KECAMATAN"])["KECAMATAN"]
        .nunique()
    )
    district_collision_count = int((district_collision > 1).sum())
    validations.append(
        {
            "VALIDASI": "Kecamatan berbeda pada OP yang sama memiliki kode berbeda",
            "STATUS": "OK" if district_collision_count == 0 else "GAGAL",
            "DETAIL": f"Benturan kode kecamatan: {district_collision_count}",
        }
    )
    same_key_count = mapping.groupby(
        ["OP", "KECAMATAN", "UJP_SIGNATURE"]
    )["KODE_ZONA"].nunique()
    inconsistent_count = int((same_key_count > 1).sum())
    validations.append(
        {
            "VALIDASI": "OP + kecamatan + pola UJP yang sama menghasilkan kode sama",
            "STATUS": "OK" if inconsistent_count == 0 else "GAGAL",
            "DETAIL": f"Kombinasi tidak konsisten: {inconsistent_count}",
        }
    )
    audit_df = pd.DataFrame(validations)
    failures = audit_df[audit_df["STATUS"] == "GAGAL"]
    if not failures.empty:
        raise ValueError("Validasi hasil gagal:\n" + failures.to_string(index=False))
    return audit_df


def find_or_create_column(ws, header_name: str, data_row_count: int) -> int:
    target = normalize_header(header_name)
    for cell in ws[1]:
        if normalize_header(cell.value) == target:
            return cell.column
    column_index = ws.max_column + 1
    source_column = max(1, column_index - 1)
    for row_index in range(1, max(ws.max_row, data_row_count + 1) + 1):
        source = ws.cell(row=row_index, column=source_column)
        target_cell = ws.cell(row=row_index, column=column_index)
        if source.has_style:
            target_cell._style = copy(source._style)
        if source.number_format:
            target_cell.number_format = source.number_format
        target_cell.alignment = copy(source.alignment)
    ws.cell(row=1, column=column_index, value=header_name)
    return column_index


def write_dataframe_sheet(wb, sheet_name: str, dataframe: pd.DataFrame) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for column_index, column_name in enumerate(dataframe.columns, start=1):
        cell = ws.cell(row=1, column=column_index, value=column_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_index, row in enumerate(dataframe.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            if isinstance(value, (tuple, list, dict)):
                value = json.dumps(value, ensure_ascii=False)
            if pd.isna(value):
                value = None
            ws.cell(row=row_index, column=column_index, value=value)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_index, column_name in enumerate(dataframe.columns, start=1):
        sample_lengths = [len(str(column_name))]
        for row_index in range(2, min(ws.max_row, 101) + 1):
            value = ws.cell(row=row_index, column=column_index).value
            if value is not None:
                sample_lengths.append(len(str(value)))
        ws.column_dimensions[get_column_letter(column_index)].width = min(
            max(sample_lengths) + 2, 45
        )


def save_output_workbook(
    file_content: bytes,
    work: pd.DataFrame,
    mapping: pd.DataFrame,
    audit_df: pd.DataFrame,
    output_file: str | Path,
) -> None:
    wb = load_workbook(io.BytesIO(file_content))
    ws = wb[SOURCE_SHEET]
    code_column_index = find_or_create_column(ws, ZONE_CODE_COLUMN, len(work))
    name_column_index = find_or_create_column(ws, ZONE_NAME_COLUMN, len(work))
    # Bersihkan nilai lama agar tidak ada sisa ketika jumlah data berkurang.
    for row_index in range(2, ws.max_row + 1):
        ws.cell(row=row_index, column=code_column_index).value = None
        ws.cell(row=row_index, column=name_column_index).value = None
    for excel_row, (_, row) in enumerate(work.iterrows(), start=2):
        ws.cell(row=excel_row, column=code_column_index, value=row[ZONE_CODE_COLUMN])
        ws.cell(row=excel_row, column=name_column_index, value=row[ZONE_NAME_COLUMN])
    mapping_output = mapping.drop(columns=["__PATTERN"], errors="ignore")
    write_dataframe_sheet(wb, MAPPING_SHEET, mapping_output)
    write_dataframe_sheet(wb, AUDIT_SHEET, audit_df)
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def upload_file() -> tuple[str, bytes] | None:
    if files is None:
        raise RuntimeError("Upload interaktif hanya tersedia saat script dijalankan di Google Colab.")
    print("Silakan upload file Excel .xlsx yang berisi sheet DATA MASTER.")
    uploaded = files.upload()
    if not uploaded:
        print("Upload dibatalkan: tidak ada file yang dipilih.")
        return None
    excel_files = [name for name in uploaded if name.lower().endswith(".xlsx")]
    if not excel_files:
        raise ValueError("File harus berformat .xlsx.")
    if len(excel_files) > 1:
        print(f"Lebih dari satu file ditemukan; file pertama yang diproses: {excel_files[0]}")
    file_name = excel_files[0]
    print(f'File "{file_name}" berhasil diupload ({len(uploaded[file_name]):,} bytes).')
    return file_name, uploaded[file_name]


def process_excel(
    file_name: str,
    file_content: bytes,
    output_path: str | Path | None = None,
) -> str:
    workbook = pd.ExcelFile(io.BytesIO(file_content))
    if SOURCE_SHEET not in workbook.sheet_names:
        raise ValueError(
            f"Sheet '{SOURCE_SHEET}' tidak ditemukan. Sheet tersedia: {workbook.sheet_names}"
        )
    df = pd.read_excel(workbook, sheet_name=SOURCE_SHEET, dtype=object)
    original_columns = list(df.columns)
    normalized_columns = [normalize_header(column) for column in original_columns]
    if len(normalized_columns) != len(set(normalized_columns)):
        duplicates = sorted(
            {column for column in normalized_columns if normalized_columns.count(column) > 1}
        )
        raise ValueError(f"Header duplikat setelah normalisasi: {duplicates}")
    df.columns = normalized_columns
    required_columns = {OP_COLUMN, KECAMATAN_COLUMN}
    missing_columns = sorted(required_columns - set(df.columns))
    if missing_columns:
        raise ValueError(f"Kolom wajib tidak ditemukan: {missing_columns}")
    if df.empty:
        raise ValueError(f"Sheet '{SOURCE_SHEET}' tidak memiliki data.")
    (
        available_ujp_columns,
        missing_standard_ujp_columns,
        auto_detected_ujp_columns,
    ) = detect_ujp_columns(list(df.columns))
    if not available_ujp_columns:
        raise ValueError(
            f"Tidak ada kolom dengan awalan '{UJP_COLUMN_PREFIX}' yang dikenali. "
            "Periksa header file."
        )
    print(f"Total baris DATA MASTER: {len(df):,}")
    print(f"Kolom UJP yang dipakai ({len(available_ujp_columns)}): {available_ujp_columns}")
    if missing_standard_ujp_columns:
        print(
            "Kolom UJP standar yang belum ada di file "
            f"({len(missing_standard_ujp_columns)}): {missing_standard_ujp_columns}"
        )
    if auto_detected_ujp_columns:
        print(
            "Kolom UJP tambahan yang terdeteksi otomatis "
            f"({len(auto_detected_ujp_columns)}): {auto_detected_ujp_columns}"
        )
    work, mapping = build_zone_mapping(df, available_ujp_columns)
    audit_df = validate_result(work, mapping)
    audit_context = pd.DataFrame(
        [
            {
                "VALIDASI": "Kolom UJP yang digunakan",
                "STATUS": "INFO",
                "DETAIL": ", ".join(available_ujp_columns),
            },
            {
                "VALIDASI": "Kolom UJP standar yang belum tersedia",
                "STATUS": "INFO",
                "DETAIL": ", ".join(missing_standard_ujp_columns) or "Tidak ada",
            },
            {
                "VALIDASI": "Kolom UJP tambahan terdeteksi otomatis",
                "STATUS": "INFO",
                "DETAIL": ", ".join(auto_detected_ujp_columns) or "Tidak ada",
            },
        ]
    )
    audit_df = pd.concat([audit_df, audit_context], ignore_index=True)
    output_file = Path(output_path or f"{Path(file_name).stem}_generate_zona.xlsx")
    save_output_workbook(file_content, work, mapping, audit_df, output_file)
    print("\n=== RINGKASAN ===")
    print(f"Total baris               : {len(work):,}")
    print(f"OP unik                   : {work['__OP'].nunique():,}")
    print(f"Kecamatan unik            : {work[['__OP', '__KECAMATAN']].drop_duplicates().shape[0]:,}")
    print(f"Kode zona unik            : {work[ZONE_CODE_COLUMN].nunique():,}")
    print(f"Panjang kode zona maksimum: {work[ZONE_CODE_COLUMN].str.len().max()}")
    print("Semua validasi            : OK")
    print("\nContoh mapping:")
    print(
        mapping[["OP", "KECAMATAN", "KODE_KECAMATAN", "SUFFIX_UJP", "KODE_ZONA", "NAMA_ZONA"]]
        .head(20)
        .to_string(index=False)
    )
    print(f"\nFile hasil: {output_file}")
    return str(output_file)




def generate(
    input_path: str | Path,
    output_path: str | Path | None = None,
) -> str:
    """Proses satu workbook lokal dan kembalikan lokasi file hasil."""
    source = Path(input_path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("File input harus berformat .xlsx")
    if not source.is_file():
        raise FileNotFoundError(f"File input tidak ditemukan: {source}")

    destination = Path(
        output_path or source.with_name(f"{source.stem}_generate_zona.xlsx")
    )
    if source.resolve() == destination.resolve():
        raise ValueError("Output harus file lain; input tidak boleh ditimpa")

    return process_excel(source.name, source.read_bytes(), destination)




def run_colab() -> None:
    """Buka dialog upload dan download hasil saat dijalankan di Colab."""
    try:
        uploaded_file = upload_file()
        if uploaded_file is None:
            return
        file_name, file_content = uploaded_file
        output_file = process_excel(file_name, file_content)
        assert files is not None
        files.download(output_file)
        print("Download file hasil dimulai.")
    except Exception as exc:
        print(f"\nERROR: {exc}")
        raise




def main() -> None:
    if files is not None:
        # Colab menyisipkan argumen internal seperti ``-f kernel.json``.
        run_colab()
        return

    if "-f" in sys.argv[1:] and "ipykernel" in sys.modules:
        raise RuntimeError(
            "Di Jupyter lokal, panggil generate('input.xlsx', 'hasil.xlsx') "
            "secara langsung."
        )

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="File .xlsx dengan sheet DATA MASTER")
    parser.add_argument("-o", "--output", help="Lokasi file hasil")
    args = parser.parse_args()
    print(generate(args.input, args.output))



if __name__ == "__main__":
    main()
