from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from scripts.penambahan_master_toko import (
    build_zone_mapping,
    canonical_ujp,
    generate,
)


def test_canonical_ujp_normalizes_common_number_formats():
    assert canonical_ujp(1000) == "N:1000"
    assert canonical_ujp("1.000") == "N:1000"
    assert canonical_ujp("1.000,00") == "N:1000"
    assert canonical_ujp(None) == "N:0"


def test_zone_mapping_is_unique_and_deterministic_for_colliding_districts():
    source = pd.DataFrame(
        [
            {"OP": "ABCDEFGH", "KECAMATAN": "Kedungjati", "UJP CDD": 100},
            {"OP": "ABCDEFGH", "KECAMATAN": "Kedungjati", "UJP CDD": 200},
            {"OP": "ABCDEFGH", "KECAMATAN": "Kedungpring", "UJP CDD": 100},
        ]
    )

    work, mapping = build_zone_mapping(source, ["UJP CDD"])

    assert mapping["KODE_ZONA"].is_unique
    assert work["KODE ZONA"].str.len().max() <= 15
    assert mapping.groupby(["OP", "KECAMATAN"])["KODE_KECAMATAN"].nunique().eq(1).all()
    assert mapping.loc[mapping["KECAMATAN"] == "KEDUNGJATI", "SUFFIX_UJP"].tolist() == [1, 2]


def test_generate_preserves_other_sheets_and_adds_mapping_and_audit(tmp_path: Path):
    input_file = tmp_path / "master.xlsx"
    output_file = tmp_path / "hasil.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "DATA MASTER"
    ws.append(["OP", "KECAMATAN", "UJP CDD", "KETERANGAN"])
    ws.append(["TGR", "Serpong", 100000, "tetap"])
    ws.append(["TGR", "Serpong", 120000, "tetap"])
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    other = wb.create_sheet("REFERENSI")
    other["A1"] = "jangan dihapus"
    wb.save(input_file)

    result = generate(input_file, output_file)

    assert result == str(output_file)
    assert output_file.exists()
    generated = load_workbook(output_file)
    assert {"DATA MASTER", "REFERENSI", "MAPPING_ZONA", "AUDIT_ZONA"} <= set(
        generated.sheetnames
    )
    assert generated["REFERENSI"]["A1"].value == "jangan dihapus"
    assert generated["DATA MASTER"]["A1"].fill.fgColor.rgb == "00FFFF00"

    result_data = pd.read_excel(output_file, sheet_name="DATA MASTER")
    assert result_data["KODE ZONA"].nunique() == 2
    assert set(result_data["NAMA ZONA"]) == {"SERPONG1", "SERPONG2"}
    audit = pd.read_excel(output_file, sheet_name="AUDIT_ZONA")
    assert "GAGAL" not in set(audit["STATUS"])


def test_generate_refuses_to_overwrite_input(tmp_path: Path):
    input_file = tmp_path / "master.xlsx"
    input_file.touch()

    with pytest.raises(ValueError, match="tidak boleh ditimpa"):
        generate(input_file, input_file)
